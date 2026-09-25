"""Parser for Cathay 00400A holdings Excel files."""

from __future__ import annotations

import math
import re
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from snapshot_validation import validate_snapshot_rows


ETF_CODE = "00400A"
SOURCE_URL = "https://www.cathaysite.com.tw/ETF/detail/EEA?tab=etf3"
SOURCE_TYPE = "data_import"
EXTRACTION_METHOD = "00400a_excel"
_STOCK_HEADER_ROW = 16
_REQUIRED_COLUMNS = ("股票代號", "股票名稱", "股數", "持股權重")
_FILENAME_RE = re.compile(
    r"^(?P<data_date>\d{4}-\d{2}-\d{2})EA\.xlsx$",
    re.IGNORECASE,
)


def parse_00400a_file(file_path: Path) -> dict:
    """Parse one 00400A workbook into canonical stock rows.

    The source date is part of Cathay's downloaded filename convention:
    YYYY-MM-DDEA.xlsx. Non-stock values in this workbook are monetary amounts,
    not percentages, so they are deliberately not imported.
    """
    file_path = Path(file_path)
    data_date = _data_date_from_filename(file_path.name)
    if data_date is None:
        return {
            "ok": False,
            "reason": "invalid_00400a_filename:expected_YYYY-MM-DDEA.xlsx",
        }
    if not file_path.is_file():
        return {"ok": False, "reason": f"file_not_found:{file_path}"}

    workbook = None
    try:
        workbook = load_workbook(file_path, read_only=True, data_only=True)
        sheet = workbook.worksheets[0]
        header_cells = next(
            sheet.iter_rows(
                min_row=_STOCK_HEADER_ROW,
                max_row=_STOCK_HEADER_ROW,
            )
        )
        columns = {
            str(cell.value).strip(): index
            for index, cell in enumerate(header_cells)
            if cell.value is not None
        }
        missing = [name for name in _REQUIRED_COLUMNS if name not in columns]
        if missing:
            return {
                "ok": False,
                "reason": f"missing_required_columns:{','.join(missing)}",
            }

        scraped_at = datetime.now().isoformat(timespec="seconds")
        stock_rows = []
        for excel_row, cells in enumerate(
            sheet.iter_rows(min_row=_STOCK_HEADER_ROW + 1),
            start=_STOCK_HEADER_ROW + 1,
        ):
            code = _stock_code(_cell_value(cells, columns["股票代號"]))
            if code is None:
                continue

            name = str(_cell_value(cells, columns["股票名稱"]) or "").strip()
            if not name:
                return {
                    "ok": False,
                    "reason": f"invalid_stock_row:{excel_row}:stock_name",
                }

            weight_cell = _cell(cells, columns["持股權重"])
            weight_pct = _weight_pct(weight_cell)
            if weight_pct is None:
                return {
                    "ok": False,
                    "reason": f"invalid_stock_row:{excel_row}:weight_pct",
                }

            shares = _shares(_cell_value(cells, columns["股數"]))
            stock_rows.append(
                {
                    "date": data_date.isoformat(),
                    "etf_code": ETF_CODE,
                    "asset_name": name,
                    "asset_type": "stock",
                    "stock_code": code,
                    "stock_name": name,
                    "shares": shares,
                    "weight_pct": weight_pct,
                    "source_url": SOURCE_URL,
                    "source_type": SOURCE_TYPE,
                    "extraction_method": EXTRACTION_METHOD,
                    "scraped_at": scraped_at,
                }
            )

        valid, reason = validate_snapshot_rows(stock_rows)
        if not valid:
            return {"ok": False, "reason": f"invalid_snapshot:{reason}"}

        return {
            "ok": True,
            "data_date": data_date,
            "stock_rows": stock_rows,
            "non_stock_rows": [],
            "source_type": SOURCE_TYPE,
        }
    except Exception as exc:
        return {
            "ok": False,
            "reason": f"invalid_00400a_workbook:{type(exc).__name__}",
        }
    finally:
        if workbook is not None:
            workbook.close()


def _data_date_from_filename(filename: str) -> date | None:
    match = _FILENAME_RE.fullmatch(filename)
    if not match:
        return None
    try:
        return date.fromisoformat(match.group("data_date"))
    except ValueError:
        return None


def _cell(cells, index):
    return cells[index] if index < len(cells) else None


def _cell_value(cells, index):
    cell = _cell(cells, index)
    return cell.value if cell is not None else None


def _stock_code(value) -> str | None:
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, int):
        text = f"{value:04d}"
    elif isinstance(value, float):
        if not value.is_integer():
            return None
        text = f"{int(value):04d}"
    else:
        text = str(value).strip()
        if re.fullmatch(r"\d+\.0", text):
            text = text[:-2]
        if text.isdigit() and len(text) < 4:
            text = text.zfill(4)

    return text if re.fullmatch(r"\d{4}", text) else None


def _shares(value):
    if value is None or str(value).strip() == "":
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        number = float(value)
    else:
        try:
            number = float(str(value).replace(",", "").strip())
        except ValueError:
            return None
    if not math.isfinite(number):
        return None
    return int(number) if number.is_integer() else number


def _weight_pct(cell) -> float | None:
    if cell is None or cell.value is None:
        return None

    value = cell.value
    try:
        if isinstance(value, str):
            text = value.strip().replace(",", "")
            if text.endswith("%"):
                number = float(text[:-1].strip())
            else:
                number = float(text)
        elif isinstance(value, bool):
            return None
        else:
            number = float(value)
            if "%" in (cell.number_format or ""):
                number *= 100.0
    except (TypeError, ValueError):
        return None

    return number if math.isfinite(number) else None
