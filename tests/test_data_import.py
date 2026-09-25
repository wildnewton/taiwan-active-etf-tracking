from datetime import date, datetime

from openpyxl import Workbook

import data_import
import db


ETF_CODE = "00400A"
DATA_DATE = date(2026, 9, 15)
STOCKS = [
    ("2301", "光寶科", 1000, "20%"),
    ("2303", "聯電", 2000, "20%"),
    ("2308", "台達電", 3000, "20%"),
    ("2317", "鴻海", 4000, "20%"),
    ("2330", "台積電", 5000, "20%"),
]


def _write_00400a_workbook(path, *, include_weight=True, stocks=STOCKS):
    workbook = Workbook()
    sheet = workbook.active
    headers = ["股票代號", "股票名稱", "股數"]
    if include_weight:
        headers.append("持股權重")
    for column, value in enumerate(headers, start=1):
        sheet.cell(row=16, column=column, value=value)

    for row_index, (stock_code, stock_name, shares, weight) in enumerate(
        stocks, start=17
    ):
        values = [stock_code, stock_name, shares]
        if include_weight:
            values.append(weight)
        for column, value in enumerate(values, start=1):
            sheet.cell(row=row_index, column=column, value=value)

    workbook.save(path)


def _existing_rows():
    scraped_at = datetime(2026, 9, 15, 21, 0).isoformat()
    return [
        {
            "date": DATA_DATE.isoformat(),
            "etf_code": ETF_CODE,
            "asset_name": f"{name}({code}.TW)",
            "asset_type": "stock",
            "stock_code": code,
            "stock_name": name,
            "shares": shares,
            "weight_pct": 20.0,
            "source_url": "https://example.test",
            "source_type": "moneydj_primary",
            "extraction_method": "test",
            "scraped_at": scraped_at,
        }
        for code, name, shares, _ in STOCKS
    ]


def test_cli_does_not_require_date(tmp_path):
    source = tmp_path / "2026-09-15EA.xlsx"
    args = data_import.build_arg_parser().parse_args(
        ["--etf", ETF_CODE, "--file", str(source)]
    )

    assert args.etf == ETF_CODE
    assert args.file == source
    assert not hasattr(args, "date")


def test_parse_00400a_derives_date_and_emits_canonical_stock_rows(tmp_path):
    source = tmp_path / "2026-09-15EA.xlsx"
    _write_00400a_workbook(source)

    result = data_import.parse_file(ETF_CODE, source)

    assert result["ok"] is True
    assert result["data_date"] == DATA_DATE
    assert result["non_stock_rows"] == []
    assert len(result["stock_rows"]) == 5

    first = result["stock_rows"][0]
    assert first["date"] == DATA_DATE.isoformat()
    assert first["etf_code"] == ETF_CODE
    assert first["asset_type"] == "stock"
    assert first["stock_code"] == "2301"
    assert first["stock_name"] == "光寶科"
    assert first["shares"] == 1000
    assert first["weight_pct"] == 20.0
    assert first["source_type"] == "data_import"
    assert first["extraction_method"] == "00400a_excel"


def test_parse_file_rejects_unsupported_etf(tmp_path):
    source = tmp_path / "2026-09-15EA.xlsx"
    _write_00400a_workbook(source)

    result = data_import.parse_file("00980A", source)

    assert result == {"ok": False, "reason": "unsupported_etf:00980A"}


def test_parse_00400a_rejects_filename_without_source_date(tmp_path):
    source = tmp_path / "holdings.xlsx"
    _write_00400a_workbook(source)

    result = data_import.parse_file(ETF_CODE, source)

    assert result == {
        "ok": False,
        "reason": "invalid_00400a_filename:expected_YYYY-MM-DDEA.xlsx",
    }


def test_malformed_00400a_workbook_fails_closed(tmp_path):
    source = tmp_path / "2026-09-15EA.xlsx"
    _write_00400a_workbook(source, include_weight=False)

    result = data_import.parse_file(ETF_CODE, source)

    assert result["ok"] is False
    assert result["reason"] == "missing_required_columns:持股權重"


def test_import_only_fills_missing_snapshot(tmp_path):
    db.init_db(tmp_path / "holdings.sqlite")
    source = tmp_path / "2026-09-15EA.xlsx"
    _write_00400a_workbook(source)

    result = data_import.import_file(ETF_CODE, source)

    assert result == {
        "ok": True,
        "etf_code": ETF_CODE,
        "data_date": DATA_DATE.isoformat(),
        "stock_rows": 5,
        "non_stock_rows": 0,
        "source_type": "data_import",
    }
    assert db.snapshot_exists(DATA_DATE, ETF_CODE) is True
    assert db.get_canonical_snapshot_source(DATA_DATE, ETF_CODE) == "data_import"

    with db._connect() as conn:
        stored = conn.execute(
            """
            SELECT stock_code, shares, weight_pct
            FROM etf_daily_holdings
            WHERE date = ? AND etf_code = ?
            ORDER BY stock_code
            """,
            (DATA_DATE.isoformat(), ETF_CODE),
        ).fetchall()

    assert stored[0] == ("2301", 1000.0, 20.0)
    assert len(stored) == 5


def test_import_refuses_to_overwrite_existing_valid_snapshot(tmp_path):
    db.init_db(tmp_path / "holdings.sqlite")
    assert db.replace_daily_snapshot(_existing_rows(), [])["inserted"] is True

    source = tmp_path / "2026-09-15EA.xlsx"
    _write_00400a_workbook(source)

    result = data_import.import_file(ETF_CODE, source)

    assert result == {
        "ok": False,
        "reason": "snapshot_already_exists",
        "etf_code": ETF_CODE,
        "data_date": DATA_DATE.isoformat(),
    }
    assert db.get_canonical_snapshot_source(DATA_DATE, ETF_CODE) == "moneydj_primary"


def test_invalid_file_never_writes_snapshot(tmp_path):
    db.init_db(tmp_path / "holdings.sqlite")
    source = tmp_path / "2026-09-15EA.xlsx"
    _write_00400a_workbook(source, include_weight=False)

    result = data_import.import_file(ETF_CODE, source)

    assert result["ok"] is False
    assert db.snapshot_exists(DATA_DATE, ETF_CODE) is False


def test_stock_like_row_with_invalid_code_fails_closed(tmp_path):
    source = tmp_path / "2026-09-15EA.xlsx"
    stocks = [
        *STOCKS,
        ("23A0", "格式異常股票", 6000, "1%"),
    ]
    _write_00400a_workbook(source, stocks=stocks)

    result = data_import.parse_file(ETF_CODE, source)

    assert result == {
        "ok": False,
        "reason": "invalid_stock_row:22:stock_code",
    }


def test_nonempty_unparseable_shares_fails_closed(tmp_path):
    source = tmp_path / "2026-09-15EA.xlsx"
    stocks = [
        *STOCKS,
        ("2603", "長榮", "not-a-number", "1%"),
    ]
    _write_00400a_workbook(source, stocks=stocks)

    result = data_import.parse_file(ETF_CODE, source)

    assert result == {
        "ok": False,
        "reason": "invalid_stock_row:22:shares",
    }
